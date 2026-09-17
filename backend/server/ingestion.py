from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook

from .config import settings
from .db import connect, init_db, json_dump, reset_content
from .text_utils import checksum, clean_text, hash_embedding, split_sentences
from .vector_backends import write_optional_vector_index


STRUCTURED_HEADERS = [
    "景区名称",
    "景点ID",
    "景点名称",
    "具体位置",
    "建筑/景观参数",
    "核心功能",
    "文化内涵",
    "详细介绍",
    "游玩亮点",
    "演艺/开放信息",
    "备注",
]


def locate_data_package() -> Path:
    package = settings.data_package_dir
    if package.exists():
        return package
    for child in settings.data_package_dir.parent.iterdir():
        if child.is_dir() and (list(child.glob("*.docx")) or list(child.glob("*.xlsx"))):
            return child
    raise FileNotFoundError(f"未找到资料包目录: {package}")


def _table_headers(table: Any) -> list[str]:
    if not table.rows:
        return []
    return [clean_text(cell.text) for cell in table.rows[0].cells]


def parse_structured_attractions(path: Path) -> list[dict[str, str]]:
    doc = Document(path)
    attractions: list[dict[str, str]] = []
    for table in doc.tables:
        headers = _table_headers(table)
        if headers[: len(STRUCTURED_HEADERS)] != STRUCTURED_HEADERS:
            continue
        for row in table.rows[1:]:
            values = [clean_text(cell.text) for cell in row.cells]
            if len(values) < len(STRUCTURED_HEADERS) or not values[1] or not values[2]:
                continue
            item = dict(zip(STRUCTURED_HEADERS, values))
            attractions.append(
                {
                    "scenic_area": item["景区名称"],
                    "spot_id": item["景点ID"],
                    "name": item["景点名称"],
                    "location": item["具体位置"],
                    "parameters": item["建筑/景观参数"],
                    "core_function": item["核心功能"],
                    "culture": item["文化内涵"],
                    "description": item["详细介绍"],
                    "highlights": item["游玩亮点"],
                    "opening": item["演艺/开放信息"],
                    "notes": item["备注"],
                    "source_file": path.name,
                }
            )
    return attractions


def parse_guide_chunks(path: Path) -> list[dict[str, Any]]:
    doc = Document(path)
    chunks: list[dict[str, Any]] = []
    current_heading = path.stem
    buffer: list[str] = []

    def flush() -> None:
        nonlocal buffer
        if not buffer:
            return
        body = "\n".join(buffer)
        for index, content in enumerate(split_sentences(body, max_chars=700)):
            chunks.append(
                {
                    "source_file": path.name,
                    "source_type": "guide",
                    "title": current_heading if index == 0 else f"{current_heading} {index + 1}",
                    "content": content,
                    "metadata": {"section": current_heading},
                }
            )
        buffer = []

    for paragraph in doc.paragraphs:
        text = clean_text(paragraph.text)
        if not text:
            continue
        style_name = getattr(paragraph.style, "name", "") or ""
        looks_like_heading = len(text) <= 28 and not text.endswith(("。", "；", "，", "：")) and "：" not in text
        if style_name.startswith("Heading") or looks_like_heading:
            flush()
            current_heading = text
        else:
            buffer.append(text)
    flush()

    for table_index, table in enumerate(doc.tables, 1):
        headers = _table_headers(table)
        for row_index, row in enumerate(table.rows[1:], 1):
            values = [clean_text(cell.text) for cell in row.cells]
            pairs = [f"{headers[i] if i < len(headers) else f'字段{i + 1}'}：{value}" for i, value in enumerate(values) if value]
            if not pairs:
                continue
            title = values[0] if values else f"表格{table_index}"
            chunks.append(
                {
                    "source_file": path.name,
                    "source_type": "guide_table",
                    "title": title,
                    "content": "；".join(pairs),
                    "metadata": {"table_index": table_index, "row_index": row_index},
                }
            )
    return chunks


def attraction_to_chunk(attraction: dict[str, str]) -> dict[str, Any]:
    fields = [
        ("景区", attraction["scenic_area"]),
        ("景点ID", attraction["spot_id"]),
        ("景点名称", attraction["name"]),
        ("具体位置", attraction["location"]),
        ("建筑/景观参数", attraction["parameters"]),
        ("核心功能", attraction["core_function"]),
        ("文化内涵", attraction["culture"]),
        ("详细介绍", attraction["description"]),
        ("游玩亮点", attraction["highlights"]),
        ("演艺/开放信息", attraction["opening"]),
        ("备注", attraction["notes"]),
    ]
    content = "\n".join(f"{key}：{value}" for key, value in fields if value)
    return {
        "source_file": attraction["source_file"],
        "source_type": "attraction",
        "title": f'{attraction["name"]}（{attraction["spot_id"]}）',
        "content": content,
        "metadata": {"spot_id": attraction["spot_id"], "scenic_area": attraction["scenic_area"]},
    }


def generate_faqs(attractions: list[dict[str, str]], guide_chunks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    faqs: list[dict[str, Any]] = []
    for item in attractions:
        source_title = f'{item["name"]}（{item["spot_id"]}）'
        candidates = [
            (f'{item["name"]}在哪里？', item["location"], ["位置", item["scenic_area"]]),
            (f'{item["name"]}有什么文化内涵？', item["culture"], ["文化", item["scenic_area"]]),
            (f'{item["name"]}适合怎么玩？', item["highlights"], ["游玩亮点", item["scenic_area"]]),
            (f'{item["name"]}的开放或演艺信息是什么？', item["opening"], ["开放信息", item["scenic_area"]]),
            (f'{item["name"]}的核心功能是什么？', item["core_function"], ["功能", item["scenic_area"]]),
        ]
        for question, answer, tags in candidates:
            if answer:
                faqs.append(
                    {
                        "question": question,
                        "answer": answer,
                        "source_title": source_title,
                        "source_file": item["source_file"],
                        "tags": tags,
                    }
                )
    for chunk in guide_chunks:
        title = chunk["title"]
        content = chunk["content"]
        if "路线规划" in content:
            faqs.append(
                {
                    "question": f"{title}怎么游览？",
                    "answer": content,
                    "source_title": title,
                    "source_file": chunk["source_file"],
                    "tags": ["路线"],
                }
            )
        if "票种" in content or "价格" in content:
            faqs.append(
                {
                    "question": f"{title}多少钱？",
                    "answer": content,
                    "source_title": title,
                    "source_file": chunk["source_file"],
                    "tags": ["票价"],
                }
            )
        if "表演时间" in content or "演出时间" in content:
            faqs.append(
                {
                    "question": f"{title}有哪些演出时间？",
                    "answer": content,
                    "source_title": title,
                    "source_file": chunk["source_file"],
                    "tags": ["演出"],
                }
            )
    return faqs


def parse_dashboard_metrics(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    headers = [clean_text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    idx = {name: i for i, name in enumerate(headers)}
    total_rows = 0
    tourists: set[str] = set()
    attraction_counts: Counter[str] = Counter()
    type_counts: Counter[str] = Counter()
    month_counts: Counter[str] = Counter()
    gender_counts: Counter[str] = Counter()
    age_buckets: Counter[str] = Counter()
    satisfaction_counts: Counter[str] = Counter()
    cost_totals: defaultdict[str, float] = defaultdict(float)
    cost_counts: defaultdict[str, int] = defaultdict(int)
    stay_by_type: defaultdict[str, list[float]] = defaultdict(list)
    matching_lingshan_rows = 0
    matching_wuxi_rows = 0

    def cell(row: tuple[Any, ...], name: str) -> Any:
        pos = idx.get(name)
        return row[pos] if pos is not None and pos < len(row) else None

    for row in sheet.iter_rows(min_row=2, values_only=True):
        total_rows += 1
        tourist_id = clean_text(cell(row, "tourist_id"))
        if tourist_id:
            tourists.add(tourist_id)
        attraction_name = clean_text(cell(row, "attraction_name"))
        attraction_type = clean_text(cell(row, "attraction_type")) or "未知类型"
        if attraction_name:
            attraction_counts[attraction_name] += 1
            if "灵山" in attraction_name or "拈花" in attraction_name:
                matching_lingshan_rows += 1
            if "无锡" in attraction_name:
                matching_wuxi_rows += 1
        type_counts[attraction_type] += 1
        visit_date = cell(row, "visit_date")
        if hasattr(visit_date, "strftime"):
            month_counts[visit_date.strftime("%Y-%m")] += 1
        gender = clean_text(cell(row, "gender")) or "未知"
        gender_counts[gender] += 1
        try:
            age = int(cell(row, "age") or 0)
        except (TypeError, ValueError):
            age = 0
        if age < 18:
            age_buckets["18岁以下"] += 1
        elif age <= 30:
            age_buckets["18-30岁"] += 1
        elif age <= 45:
            age_buckets["31-45岁"] += 1
        elif age <= 60:
            age_buckets["46-60岁"] += 1
        else:
            age_buckets["60岁以上"] += 1
        satisfaction = clean_text(cell(row, "satisfaction")) or "未评价"
        satisfaction_counts[satisfaction] += 1
        for field in ["ticket_cost", "food_cost", "shopping_cost", "transport_cost", "entertainment_cost", "total_cost"]:
            try:
                value = float(cell(row, field) or 0)
            except (TypeError, ValueError):
                value = 0.0
            cost_totals[field] += value
            cost_counts[field] += 1
        try:
            stay = float(cell(row, "stay_duration") or 0)
        except (TypeError, ValueError):
            stay = 0.0
        if stay > 0:
            stay_by_type[attraction_type].append(stay)

    workbook.close()
    avg_costs = {field: round(cost_totals[field] / max(1, cost_counts[field]), 2) for field in cost_totals}
    avg_stay_by_type = {
        key: round(sum(values) / len(values), 2)
        for key, values in stay_by_type.items()
        if values
    }
    return {
        "source_file": path.name,
        "total_rows": total_rows,
        "unique_tourists": len(tourists),
        "matching_lingshan_rows": matching_lingshan_rows,
        "matching_wuxi_rows": matching_wuxi_rows,
        "data_scope_note": "Excel 为长三角多景区游客行为样本。大屏用于观察人群偏好、季节趋势和消费结构；景点讲解仍以本地资料包与景区现场公告为准。",
        "top_attractions": attraction_counts.most_common(10),
        "attraction_types": type_counts.most_common(),
        "monthly_visits": sorted(month_counts.items()),
        "gender_distribution": gender_counts.most_common(),
        "age_buckets": age_buckets.most_common(),
        "satisfaction_distribution": satisfaction_counts.most_common(),
        "average_costs": avg_costs,
        "average_stay_by_type": sorted(avg_stay_by_type.items(), key=lambda item: item[1], reverse=True)[:10],
    }


def rebuild_knowledge() -> dict[str, Any]:
    package = locate_data_package()
    docx_files = sorted(package.glob("*.docx"))
    xlsx_files = sorted(package.glob("*.xlsx"))
    structured_file = next((path for path in docx_files if len(Document(path).tables) >= 2 and len([p for p in Document(path).paragraphs if clean_text(p.text)]) < 20), None)
    guide_files = [path for path in docx_files if path != structured_file]
    if structured_file is None:
        raise FileNotFoundError("资料包中未找到景点结构化数据集 docx")

    attractions = parse_structured_attractions(structured_file)
    guide_chunks: list[dict[str, Any]] = []
    for guide in guide_files:
        guide_chunks.extend(parse_guide_chunks(guide))
    chunks = [attraction_to_chunk(item) for item in attractions] + guide_chunks
    faqs = generate_faqs(attractions, guide_chunks)
    metrics = parse_dashboard_metrics(xlsx_files[0]) if xlsx_files else {}
    vector_result = write_optional_vector_index(chunks)

    now = datetime.now().isoformat(timespec="seconds")
    conn = connect()
    init_db(conn)
    reset_content(conn)
    try:
        for file in docx_files + xlsx_files:
            records_count = metrics.get("total_rows", 0) if file.suffix.lower() == ".xlsx" else 0
            conn.execute(
                "INSERT INTO source_files(name,path,file_type,checksum,imported_at,records_count) VALUES(?,?,?,?,?,?)",
                (file.name, str(file), file.suffix.lower().lstrip("."), checksum(file), now, records_count),
            )
        conn.executemany(
            """
            INSERT INTO attractions(
              scenic_area,spot_id,name,location,parameters,core_function,culture,description,highlights,opening,notes,source_file
            ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)
            """,
            [
                (
                    item["scenic_area"],
                    item["spot_id"],
                    item["name"],
                    item["location"],
                    item["parameters"],
                    item["core_function"],
                    item["culture"],
                    item["description"],
                    item["highlights"],
                    item["opening"],
                    item["notes"],
                    item["source_file"],
                )
                for item in attractions
            ],
        )
        for chunk in chunks:
            embedding = hash_embedding(f'{chunk["title"]}\n{chunk["content"]}')
            cursor = conn.execute(
                "INSERT INTO knowledge_chunks(source_file,source_type,title,content,metadata_json,embedding_json) VALUES(?,?,?,?,?,?)",
                (
                    chunk["source_file"],
                    chunk["source_type"],
                    chunk["title"],
                    chunk["content"],
                    json_dump(chunk["metadata"]),
                    json_dump(embedding),
                ),
            )
            chunk_id = cursor.lastrowid
            conn.execute(
                "INSERT INTO knowledge_fts(title,content,chunk_id) VALUES(?,?,?)",
                (chunk["title"], chunk["content"], chunk_id),
            )
        conn.executemany(
            "INSERT INTO faqs(question,answer,source_title,source_file,tags_json) VALUES(?,?,?,?,?)",
            [(faq["question"], faq["answer"], faq["source_title"], faq["source_file"], json_dump(faq["tags"])) for faq in faqs],
        )
        conn.execute(
            "INSERT INTO dashboard_metrics(id,metrics_json,updated_at) VALUES(1,?,?)",
            (json_dump(metrics), now),
        )
        conn.execute(
            """
            INSERT OR IGNORE INTO digital_human_config(
              id,name,voice,persona,avatar_mode,rtx_enabled,opentalking_base_url,audio2face_base_url,updated_at
            ) VALUES(1,?,?,?,?,?,?,?,?)
            """,
            (
                "灵境导游",
                settings.vivo_tts_voice,
                "温和、准确、只依据资料包回答的景区数字人导游",
                "opentalking" if settings.opentalking_base_url else "2d-fallback",
                0,
                settings.opentalking_base_url,
                "",
                now,
            ),
        )
        conn.commit()
    finally:
        conn.close()

    export_dir = settings.storage_dir / "exports"
    export_dir.mkdir(parents=True, exist_ok=True)
    (export_dir / "faq_test_set.json").write_text(json_dump(faqs), encoding="utf-8")
    (export_dir / "dashboard_metrics.json").write_text(json_dump(metrics), encoding="utf-8")
    (export_dir / "knowledge_manifest.json").write_text(
        json_dump(
            {
                "package": str(package),
                "source_files": [path.name for path in docx_files + xlsx_files],
                "attractions": len(attractions),
                "chunks": len(chunks),
                "faqs": len(faqs),
                "vector_backend": vector_result,
                "rebuilt_at": now,
            }
        ),
        encoding="utf-8",
    )
    return {
        "package": str(package),
        "sources": [path.name for path in docx_files + xlsx_files],
        "attractions": len(attractions),
        "chunks": len(chunks),
        "faqs": len(faqs),
        "dashboard_rows": metrics.get("total_rows", 0),
        "matching_lingshan_rows": metrics.get("matching_lingshan_rows", 0),
        "matching_wuxi_rows": metrics.get("matching_wuxi_rows", 0),
        "vector_backend": vector_result,
        "exports": str(export_dir),
        "rebuilt_at": now,
    }
