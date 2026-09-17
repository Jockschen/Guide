from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from pptx import Presentation


def main() -> int:
    pptx_path = Path(sys.argv[1]).resolve()
    if not pptx_path.is_file():
        raise AssertionError(f"PPTX not found: {pptx_path}")
    if pptx_path.stat().st_size <= 0:
        raise AssertionError("PPTX is empty")

    prs = Presentation(str(pptx_path))
    if len(prs.slides) != 10:
        raise AssertionError(f"Expected 10 slides, found {len(prs.slides)}")
    ratio = prs.slide_width / prs.slide_height
    if abs(ratio - (16 / 9)) > 0.001:
        raise AssertionError(f"Unexpected aspect ratio: {ratio:.6f}")

    with zipfile.ZipFile(pptx_path) as archive:
        names = archive.namelist()
        slide_xml = sorted(
            name for name in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
        )
        notes_xml = sorted(
            name for name in names if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", name)
        )
        media = sorted(name for name in names if name.startswith("ppt/media/"))

        if len(slide_xml) != 10:
            raise AssertionError(f"Expected 10 slide XML files, found {len(slide_xml)}")
        if len(notes_xml) != 10:
            raise AssertionError(f"Expected 10 notes slides, found {len(notes_xml)}")
        if len(media) < 10:
            raise AssertionError(f"Expected at least 10 media files, found {len(media)}")

        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        nonempty_notes = 0
        for name in notes_xml:
            root = ET.fromstring(archive.read(name))
            note_text = "".join(node.text or "" for node in root.findall(".//a:t", ns)).strip()
            if note_text:
                nonempty_notes += 1
        if nonempty_notes != 10:
            raise AssertionError(f"Expected 10 non-empty notes, found {nonempty_notes}")

    print(f"file={pptx_path}")
    print(f"size_bytes={pptx_path.stat().st_size}")
    print(f"slides={len(prs.slides)}")
    print(f"aspect_ratio={ratio:.6f}")
    print(f"media_files={len(media)}")
    print(f"notes={len(notes_xml)}; nonempty_notes={nonempty_notes}")
    print("verification=PASS")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
